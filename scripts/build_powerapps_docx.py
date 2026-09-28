"""Converte os .md do Power Apps em um único Word (.docx) bem formatado."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DHL_YELLOW = RGBColor(0xB8, 0x8A, 0x00)
DHL_RED = RGBColor(0xD4, 0x05, 0x11)
CODE_BG = "F2F3F5"

SRC_FILES = [
    "/app/memory/PowerApps_CopiaEcola_PowerFx.md",
    "/app/memory/PowerApps_Blueprint_SmartTime.md",
]
OUT = "/app/frontend/public/SmartTime_PowerApps_Formulas.docx"


def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        shade(p, CODE_BG)
        run = p.add_run(ln if ln else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        r = run._element
        r.rPr.rFonts.set(qn("w:cs"), "Consolas")


def strip_md(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


def add_para_with_bold(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    # handle **bold** and `code`
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9.5)
        else:
            p.add_run(part)
    return p


def add_table(doc, rows):
    cols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            para = c.paragraphs[0]
            run = para.add_run(strip_md(cell.strip()))
            run.font.size = Pt(8.5)
            if i == 0:
                run.bold = True
    return t


def convert(doc, path, first):
    with open(path, encoding="utf-8") as f:
        raw = f.read().split("\n")

    i = 0
    in_code = False
    code_lines = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            # drop separator row (---)
            rows = [r for r in table_buf if not re.match(r"^\s*\|?[\s:\-\|]+\|?\s*$", r)]
            parsed = []
            for r in rows:
                cells = [c for c in r.strip().strip("|").split("|")]
                parsed.append(cells)
            if parsed:
                add_table(doc, parsed)
            doc.add_paragraph()
            table_buf = []

    for line in raw:
        if line.strip().startswith("```"):
            if in_code:
                add_code(doc, code_lines)
                doc.add_paragraph()
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.strip().startswith("|"):
            table_buf.append(line)
            continue
        else:
            flush_table()

        s = line.rstrip()
        if not s.strip():
            continue
        if s.startswith("# "):
            if not first:
                doc.add_page_break()
            h = doc.add_heading(strip_md(s[2:]), level=0)
            first = False
        elif s.startswith("## "):
            doc.add_heading(strip_md(s[3:]), level=1)
        elif s.startswith("### "):
            doc.add_heading(strip_md(s[4:]), level=2)
        elif s.startswith(">"):
            p = add_para_with_bold(doc, s.lstrip("> ").strip())
            p.paragraph_format.left_indent = Inches(0.2)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif re.match(r"^\s*[-*] ", s):
            add_para_with_bold(doc, re.sub(r"^\s*[-*] ", "", s), style="List Bullet")
        elif re.match(r"^\s*\d+\. ", s):
            add_para_with_bold(doc, re.sub(r"^\s*\d+\. ", "", s), style="List Number")
        elif set(s.strip()) <= set("-"):
            continue
        else:
            add_para_with_bold(doc, s)
    flush_table()
    return first


doc = Document()
# base styles
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)

# Cover
title = doc.add_heading("SmartTime!", level=0)
sub = doc.add_paragraph("Fórmulas Power Fx — Esqueleto Copia-e-Cola + Blueprint Completo")
sub.runs[0].bold = True
sub.runs[0].font.size = Pt(14)
sub.runs[0].font.color.rgb = DHL_RED
info = doc.add_paragraph("Recriação do app de Gestão de Horas Extras (DHL) como Canvas App no Microsoft Power Apps.\n"
                         "Documento 1: fórmulas prontas para colar por controle.  Documento 2: blueprint (tabelas, telas, fluxos).")
info.runs[0].font.size = Pt(10)
doc.add_paragraph()

first = True
for path in SRC_FILES:
    first = convert(doc, path, first)

doc.save(OUT)
print("Salvo em", OUT)
